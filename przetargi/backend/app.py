from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io
import os
import re

app = Flask(__name__)
CORS(app)

def format_date(date_str):
    """Konwertuje datę z formatu YYYY-MM-DD na DD.MM.YYYY"""
    if not date_str:
        return ''
    try:
        parts = date_str.split('-')
        if len(parts) == 3:
            return f"{parts[2]}.{parts[1]}.{parts[0]}"
    except:
        pass
    return date_str

def replace_in_paragraph(paragraph, old_text, new_text):
    """Zamienia tekst w paragrafie zachowując formatowanie"""
    if old_text in paragraph.text:
        inline = paragraph.runs
        for run in inline:
            if old_text in run.text:
                run.text = run.text.replace(old_text, new_text)

def replace_in_cell(cell, replacements):
    """Zamienia tekst w komórce tabeli"""
    for paragraph in cell.paragraphs:
        for old_text, new_text in replacements.items():
            if old_text in paragraph.text:
                # Próbujemy zastąpić w runach
                for run in paragraph.runs:
                    for old, new in replacements.items():
                        if old in run.text:
                            run.text = run.text.replace(old, new if new else '..................')

def fill_template(template_path, data):
    """Wypełnia szablon dokumentu danymi z formularza"""
    doc = Document(template_path)
    
    # Mapowanie pól formularza na miejsca w dokumencie
    replacements = {
        # Oznaczenie sprawy (w nagłówkach)
        '...........................': data.get('oznaczenieSpawy', ''),
        
        # Zamawiający
        '.................................................': data.get('nazwaZamawiajacego', ''),
        
        # Przedmiot zamówienia
        '.........................................': data.get('nazwaPrzedmiotu', ''),
        '............................': data.get('nazwaPrzedmiotu', ''),
        
        # Wartość
        '.........................': data.get('wartoscZamowienia', ''),
        '......................': data.get('wartoscEuro', ''),
        
        # Data ogłoszenia BZP
        '.............................. r.': format_date(data.get('dataOgloszeniaBZP', '')) + ' r.' if data.get('dataOgloszeniaBZP') else '.............................. r.',
        
        # Numer ogłoszenia
        '.................': data.get('numerOgloszeniaBZP', ''),
        
        # Termin składania ofert
        '................ ...............': format_date(data.get('terminSkladaniaOfertData', '')),
        '....... : .': data.get('terminSkladaniaOfertGodzina', '').replace(':', ':') if data.get('terminSkladaniaOfertGodzina') else '....... : .',
        
        # Data zawarcia umowy
        '............... r.': format_date(data.get('dataZawarciaUmowy', '')) + ' r.' if data.get('dataZawarciaUmowy') else '............... r.',
        
        # Wykonawca umowy
        '.......................................': data.get('wykonawcaUmowy', ''),
        '......................................': data.get('wykonawcaUmowy', ''),
        
        # Osoba sporządzająca
        '......................................................................................': data.get('osobaSPorzadzajaca', ''),
    }
    
    # Zamiana w paragrafach
    for paragraph in doc.paragraphs:
        full_text = paragraph.text
        for old_text, new_text in replacements.items():
            if old_text in full_text and new_text:
                for run in paragraph.runs:
                    if old_text in run.text:
                        run.text = run.text.replace(old_text, new_text)
    
    # Zamiana w tabelach
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text
                for old_text, new_text in replacements.items():
                    if old_text in cell_text and new_text:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                if old_text in run.text:
                                    run.text = run.text.replace(old_text, new_text)
    
    # Specjalne wypełnienie dla konkretnych tabel
    # Tabela 0 - Zamawiający i przedmiot
    if len(doc.tables) > 0:
        table = doc.tables[0]
        # Wiersz 0 - Zamawiający
        if data.get('nazwaZamawiajacego'):
            cell = table.rows[0].cells[1]
            for para in cell.paragraphs:
                if '.................................................' in para.text:
                    for run in para.runs:
                        run.text = run.text.replace('.................................................', data['nazwaZamawiajacego'])
        
        # Wiersz 1 - Przedmiot zamówienia
        if data.get('nazwaPrzedmiotu'):
            cell = table.rows[1].cells[1]
            for para in cell.paragraphs:
                for run in para.runs:
                    if '.........................................' in run.text:
                        run.text = run.text.replace('.........................................', data['nazwaPrzedmiotu'])
        
        # Wiersz 2 - Wartość
        if data.get('wartoscZamowienia'):
            cell = table.rows[2].cells[1]
            for para in cell.paragraphs:
                for run in para.runs:
                    if '.........................' in run.text:
                        run.text = run.text.replace('.........................', data['wartoscZamowienia'])
                    if '......................' in run.text:
                        run.text = run.text.replace('......................', data.get('wartoscEuro', ''))
    
    # Tabela 3 - Ogłoszenie o zamówieniu
    if len(doc.tables) > 3:
        table = doc.tables[3]
        # Wiersz 1 - Ogłoszenie BZP
        if data.get('dataOgloszeniaBZP') or data.get('numerOgloszeniaBZP'):
            cell = table.rows[1].cells[1]
            for para in cell.paragraphs:
                for run in para.runs:
                    if '.............................. r.' in run.text:
                        run.text = run.text.replace('.............................. r.', format_date(data.get('dataOgloszeniaBZP', '')) + ' r.')
                    if '.................' in run.text:
                        run.text = run.text.replace('.................', data.get('numerOgloszeniaBZP', ''))
        
        # Wiersz 3 - SWZ adres
        if data.get('adresSWZ'):
            cell = table.rows[3].cells[1]
            for para in cell.paragraphs:
                for run in para.runs:
                    if '......................' in run.text:
                        run.text = run.text.replace('......................', data['adresSWZ'])
    
    # Tabela 4 - Termin składania i otwarcie ofert
    if len(doc.tables) > 4:
        table = doc.tables[4]
        # SWZ adres
        if data.get('adresSWZ'):
            cell = table.rows[0].cells[1]
            for para in cell.paragraphs:
                for run in para.runs:
                    if '......................' in run.text:
                        run.text = run.text.replace('......................', data['adresSWZ'])
        
        # Termin składania ofert
        if data.get('terminSkladaniaOfertData'):
            cell = table.rows[1].cells[1]
            for para in cell.paragraphs:
                for run in para.runs:
                    if '................' in run.text:
                        run.text = run.text.replace('................', format_date(data['terminSkladaniaOfertData']))
                    if '....... : .' in run.text:
                        run.text = run.text.replace('....... : .', data.get('terminSkladaniaOfertGodzina', ''))
        
        # Otwarcie ofert
        if data.get('dataOtwarciaOfert'):
            cell = table.rows[2].cells[1]
            for para in cell.paragraphs:
                for run in para.runs:
                    if '................' in run.text:
                        run.text = run.text.replace('................', format_date(data['dataOtwarciaOfert']))
    
    # Tabela 8 - Udzielenie zamówienia
    if len(doc.tables) > 8:
        table = doc.tables[8]
        # Wiersz 4 - Udzielenie zamówienia
        if data.get('dataZawarciaUmowy') or data.get('wykonawcaUmowy') or data.get('kwotaUmowy'):
            cell = table.rows[4].cells[1]
            for para in cell.paragraphs:
                for run in para.runs:
                    if '............... r.' in run.text:
                        run.text = run.text.replace('............... r.', format_date(data.get('dataZawarciaUmowy', '')) + ' r.')
                    if '.....................................' in run.text:
                        run.text = run.text.replace('.....................................', data.get('wykonawcaUmowy', ''))
    
    # Tabela 9 - Ogłoszenie o wyniku i osoba sporządzająca
    if len(doc.tables) > 9:
        table = doc.tables[9]
        # Wiersz 0 - Ogłoszenie o wyniku
        if data.get('dataOgloszeniaWyniku') or data.get('numerOgloszeniaWyniku'):
            cell = table.rows[0].cells[1]
            for para in cell.paragraphs:
                for run in para.runs:
                    if '................................ r.' in run.text:
                        run.text = run.text.replace('................................ r.', format_date(data.get('dataOgloszeniaWyniku', '')) + ' r.')
                    if '.................' in run.text:
                        run.text = run.text.replace('.................', data.get('numerOgloszeniaWyniku', ''))
        
        # Wiersz 3 - Osoba sporządzająca
        if data.get('osobaSPorzadzajaca'):
            cell = table.rows[3].cells[1]
            for para in cell.paragraphs:
                for run in para.runs:
                    if '......................' in run.text:
                        run.text = run.text.replace('......................', data['osobaSPorzadzajaca'])
        
        # Wiersz 4 - Zatwierdzenie
        if data.get('osobaZatwierdzajaca'):
            cell = table.rows[4].cells[1]
            for para in cell.paragraphs:
                for run in para.runs:
                    if '...........................................' in run.text:
                        run.text = run.text.replace('...........................................', data['osobaZatwierdzajaca'])
    
    return doc

@app.route('/api/generate-doc', methods=['POST'])
def generate_doc():
    try:
        data = request.json
        
        # Ścieżka do szablonu
        template_path = os.path.join(os.path.dirname(__file__), 'szablon_protokol.docx')
        
        if not os.path.exists(template_path):
            return jsonify({'error': 'Szablon nie został znaleziony'}), 404
        
        # Wypełnij szablon
        doc = fill_template(template_path, data)
        
        # Zapisz do bufora
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        
        return send_file(
            file_stream,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name='Protokol_Przetargu.docx'
        )
        
    except Exception as e:
        print(f"Error: {str(e)}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/health', methods=['GET'])
def health_check():
    return jsonify({'status': 'ok'})

if __name__ == '__main__':
    app.run(debug=True, port=5000)
